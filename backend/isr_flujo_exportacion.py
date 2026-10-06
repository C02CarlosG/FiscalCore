"""Excel del ISR por flujo: una hoja con los renglones que componen una cifra y una hoja con el resumen del mes y del
acumulado del ejercicio. Mismas cifras que la pantalla."""
from __future__ import annotations

from io import BytesIO

import openpyxl
from openpyxl.utils import get_column_letter

from .iva_flujo_exportacion import FORMATO_MONEDA, _NEGRITA, _asignar, _encabezados

_COLUMNAS = [
    ("Fecha de emisión", lambda r: r["fecha_emision"], "fecha"),
    ("Fecha de efecto", lambda r: r["fecha_efecto"], "fecha"),
    ("UUID", lambda r: r["uuid"], "texto"),
    ("UUID del REP", lambda r: r["uuid_pago"], "texto"),
    ("Tipo", lambda r: r["tipo_comprobante"], "texto"),
    ("Origen", lambda r: r["origen"], "texto"),
    ("RFC contraparte", lambda r: r["contraparte_rfc"], "texto"),
    ("Contraparte", lambda r: r["contraparte"], "texto"),
    ("Base", lambda r: r["base"], "moneda"),
    ("ISR retenido", lambda r: r["retencion"], "moneda"),
    ("Gravado (nómina)", lambda r: (r["nomina"] or {}).get("gravado"), "moneda"),
    ("Exento (nómina)", lambda r: (r["nomina"] or {}).get("exento"), "moneda"),
    ("Marcas", lambda r: ", ".join(r["marcas"]), "texto"),
    ("Motivo", lambda r: r["motivo"], "texto"),
]


def construir(periodo: str, lado: str, bloque: str, filas: list[dict], resumen: dict, regimen: dict | None = None) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Detalle"
    _encabezados(ws, [c[0] for c in _COLUMNAS])
    anchos = [max(len(c[0]), 10) for c in _COLUMNAS]
    for r in filas:
        ws.append([None] * len(_COLUMNAS))
        for i, (_, obtener, tipo) in enumerate(_COLUMNAS, start=1):
            valor = obtener(r)
            _asignar(ws.cell(row=ws.max_row, column=i), valor, tipo)
            if isinstance(valor, str):
                anchos[i - 1] = min(max(anchos[i - 1], len(valor)), 60)
    for i, ancho in enumerate(anchos, start=1):
        ws.column_dimensions[get_column_letter(i)].width = ancho + 2

    wr = wb.create_sheet("Resumen")
    _encabezados(wr, [f"ISR por flujo {periodo}", "Mes", "Acumulado"])

    def _fila(etiqueta: str, obtener, negrita: bool = False) -> None:
        wr.append([etiqueta, float(obtener(resumen["mes"])), float(obtener(resumen["acumulado"]))])
        for celda in wr[wr.max_row][1:]:
            celda.number_format = FORMATO_MONEDA
        if negrita:
            wr[wr.max_row][0].font = _NEGRITA

    _fila("Ingresos de contado", lambda b: b["ingresos"]["contado"])
    _fila("Ingresos cobrados a crédito", lambda b: b["ingresos"]["credito"])
    _fila("Devoluciones y descuentos", lambda b: b["ingresos"]["devoluciones"])
    _fila("Total de ingresos", lambda b: b["ingresos"]["total"], negrita=True)
    _fila("Compras y gastos", lambda b: b["deducciones"]["compras_y_gastos"])
    _fila("Devoluciones recibidas (ya restadas)", lambda b: b["deducciones"]["devoluciones_recibidas"])
    _fila("Nómina deducible", lambda b: b["deducciones"]["nomina"]["deducible"])
    _fila("Total de deducciones", lambda b: b["deducciones"]["total"], negrita=True)
    _fila("Utilidad fiscal estimada", lambda b: b["utilidad_fiscal_estimada"], negrita=True)
    wr.append([])
    _fila("ISR retenido a favor", lambda b: b["ingresos"]["retenciones_a_favor"])
    _fila("ISR retenido a cargo (trabajadores)", lambda b: b["retenciones_a_cargo"]["trabajadores"])
    _fila("ISR retenido a cargo (proveedores)", lambda b: b["retenciones_a_cargo"]["proveedores"])
    wr.append([])
    _fila("Inversiones identificadas (no suman)", lambda b: b["deducciones"]["inversiones"]["base"])
    _fila("Ingresos no considerados (base)", lambda b: b["ingresos"]["no_considerados"]["base"])
    _fila("Deducciones no consideradas (base)", lambda b: b["deducciones"]["no_considerados"]["base"])
    wr.append([])
    wr.append(["Nómina exenta deducible", float(resumen["porcentaje_nomina_exenta"])])
    wr.cell(row=wr.max_row, column=2).number_format = "0%"
    if regimen:
        wr.append(["Régimen", f'{regimen.get("codigo") or "sin régimen"} ({regimen.get("modulo")})'])
        for aviso in regimen.get("avisos", []):
            wr.append([aviso])
    for a in resumen.get("advertencias", []):
        wr.append([a["mensaje"]])
    wr.append([])
    wr.append(["Estimación del flujo: no calcula el pago provisional."])
    wr.column_dimensions["A"].width = 38
    wr.column_dimensions["B"].width = 18
    wr.column_dimensions["C"].width = 18

    salida = BytesIO()
    wb.save(salida)
    return salida.getvalue()
