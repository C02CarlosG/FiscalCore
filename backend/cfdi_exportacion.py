"""Excel del listado de CFDI: una hoja con las filas y las columnas visibles, y una de
totales (periodo y acumulado) con las mismas cifras que la pantalla."""
from __future__ import annotations

from datetime import date, datetime
from io import BytesIO
from typing import Any

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .cfdi_columnas import Columna

FORMATO_MONEDA = '"$"#,##0.00'
FORMATO_FECHA = "dd/mm/yyyy"
FORMATO_FECHA_HORA = "dd/mm/yyyy hh:mm"

_TOTALES = (
    ("conteo", "CFDI"), ("retencion_iva", "Retención IVA"), ("retencion_ieps", "Retención IEPS"),
    ("retencion_isr", "Retención ISR"), ("traslado_iva", "Traslado IVA"), ("traslado_ieps", "Traslado IEPS"),
    ("traslado_isr", "Traslado ISR"), ("total_retenciones", "Total retenciones"), ("subtotal", "Subtotal"),
    ("descuento", "Descuento"), ("neto", "Neto"), ("total", "Total"),
)
_ENCABEZADO = Font(bold=True, color="FFFFFF")
_RELLENO = PatternFill("solid", fgColor="334155")


def _celda(col: Columna, valor: Any) -> tuple[Any, str | None]:
    """Valor para la celda y su formato numérico. Vacío = celda vacía."""
    if valor is None or valor == "":
        return None, None
    if col.tipo_dato == "moneda":
        return float(valor), FORMATO_MONEDA
    if col.tipo_dato == "fecha":
        return date.fromisoformat(str(valor)[:10]), FORMATO_FECHA
    if col.tipo_dato == "fecha_hora":
        return datetime.fromisoformat(str(valor)), FORMATO_FECHA_HORA
    if col.tipo_dato == "booleano":
        return "Sí" if valor else "No", None
    if col.tipo_dato == "lista":
        return ", ".join(str(v) for v in valor) if isinstance(valor, list) else str(valor), None
    if col.tipo_dato == "numero":
        return float(valor), None
    return str(valor), None


def _encabezados(ws, textos: list[str]) -> None:
    ws.append(textos)
    for celda in ws[1]:
        celda.font = _ENCABEZADO
        celda.fill = _RELLENO
        celda.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"


def construir(cols: list[Columna], items: list[dict], resumen: dict) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "CFDI"
    _encabezados(ws, [c.etiqueta for c in cols])
    anchos = [max(len(c.etiqueta), 10) for c in cols]
    for item in items:
        fila = []
        for i, col in enumerate(cols):
            valor, formato = _celda(col, item.get(col.clave))
            fila.append((valor, formato))
            if isinstance(valor, str):
                anchos[i] = min(max(anchos[i], len(valor)), 60)
        ws.append([v for v, _ in fila])
        for celda, (_, formato) in zip(ws[ws.max_row], fila):
            if formato:
                celda.number_format = formato
    for i, ancho in enumerate(anchos, start=1):
        ws.column_dimensions[get_column_letter(i)].width = ancho + 2

    wt = wb.create_sheet("Totales")
    _encabezados(wt, ["", *[etiqueta for _, etiqueta in _TOTALES]])
    for nombre, bloque in (("Periodo", resumen["totales"]["periodo"]), ("Acumulado", resumen["totales"]["acumulado"])):
        wt.append([nombre, *[bloque.get(clave) for clave, _ in _TOTALES]])
        for celda, (clave, _) in zip(wt[wt.max_row][1:], _TOTALES):
            if clave != "conteo":
                celda.number_format = FORMATO_MONEDA
    wt.column_dimensions["A"].width = 14
    for i in range(2, len(_TOTALES) + 2):
        wt.column_dimensions[get_column_letter(i)].width = 18

    salida = BytesIO()
    wb.save(salida)
    return salida.getvalue()
