"""Excel del IVA por flujo: una hoja con los renglones que componen una cifra (con las bases
e IVA por tasa) y una hoja con el resumen del periodo. Mismas cifras que la pantalla."""
from __future__ import annotations

from datetime import date, datetime
from io import BytesIO
from typing import Any, Optional

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

FORMATO_MONEDA = '"$"#,##0.00'
FORMATO_FECHA = "dd/mm/yyyy"

_ENCABEZADO = Font(bold=True, color="FFFFFF")
_RELLENO = PatternFill("solid", fgColor="334155")
_NEGRITA = Font(bold=True)

# (encabezado, clave | función del renglón, tipo)
_COLUMNAS = [
    ("Fecha de emisión", lambda r: r["fecha_emision"], "fecha"),
    ("Fecha de pago", lambda r: r["fecha_pago"], "fecha"),
    ("UUID", lambda r: r["uuid"], "texto"),
    ("UUID del REP", lambda r: r["uuid_pago"], "texto"),
    ("Parcialidad", lambda r: r["parcialidad"], "numero"),
    ("Tipo", lambda r: r["tipo_comprobante"], "texto"),
    ("RFC contraparte", lambda r: r["contraparte_rfc"], "texto"),
    ("Contraparte", lambda r: r["contraparte"], "texto"),
    ("Base 16 %", lambda r: r["bases"]["16"], "moneda"),
    ("IVA 16 %", lambda r: r["iva"]["16"], "moneda"),
    ("Base 8 %", lambda r: r["bases"]["8"], "moneda"),
    ("IVA 8 %", lambda r: r["iva"]["8"], "moneda"),
    ("Base 0 %", lambda r: r["bases"]["0"], "moneda"),
    ("Base exento", lambda r: r["bases"]["exento"], "moneda"),
    ("Base otras tasas", lambda r: r["bases"]["otras"], "moneda"),
    ("IVA otras tasas", lambda r: r["iva"]["otras"], "moneda"),
    ("Base no objeto", lambda r: r["bases"]["no_objeto"], "moneda"),
    ("IVA total", lambda r: r["iva_total"], "moneda"),
    ("Retención de IVA", lambda r: r["retencion"], "moneda"),
    ("Total del documento", lambda r: r["total_documento"], "moneda"),
    ("Marcas", lambda r: ", ".join(r["marcas"]), "texto"),
    ("Motivo", lambda r: r["motivo"], "texto"),
    ("Ajuste", lambda r: _ajuste(r["ajuste"]), "texto"),
]

_ETIQUETAS_ORIGEN = {
    "contado": "Contado",
    "credito": "Cobro o pago de crédito",
    "notas_credito": "Notas de crédito",
}


def _ajuste(ajuste: Optional[dict]) -> Optional[str]:
    if not ajuste:
        return None
    destino = f" → {ajuste['periodo_destino']}" if ajuste.get("periodo_destino") else ""
    return f"{ajuste['accion']}{destino}: {ajuste.get('motivo') or ''}".strip()


def _asignar(celda, valor: Any, tipo: str) -> None:
    """Valor y formato de una celda. Un texto nunca se interpreta como fórmula."""
    if valor is None or valor == "":
        return
    if tipo == "fecha":
        celda.value = datetime.fromisoformat(str(valor)) if "T" in str(valor) else datetime.combine(
            date.fromisoformat(str(valor)[:10]), datetime.min.time())
        celda.number_format = FORMATO_FECHA
    elif tipo == "moneda":
        celda.value = float(valor)
        celda.number_format = FORMATO_MONEDA
    elif tipo == "numero":
        celda.value = float(valor) if not isinstance(valor, int) else valor
    else:
        celda.value = str(valor)
        celda.data_type = "s"            # "=..." se queda como texto: no se ejecuta como fórmula


def _encabezados(ws, textos: list[str]) -> None:
    ws.append(textos)
    for celda in ws[1]:
        celda.font = _ENCABEZADO
        celda.fill = _RELLENO
        celda.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"


def construir(periodo: str, direccion: str, origen: str, filas: list[dict], resumen: dict) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Detalle"
    _encabezados(ws, [c[0] for c in _COLUMNAS])
    anchos = [max(len(c[0]), 10) for c in _COLUMNAS]
    for r in filas:
        ws.append([None] * len(_COLUMNAS))
        for i, (_, obtener, tipo) in enumerate(_COLUMNAS, start=1):
            valor = obtener(r)
            _asignar(ws.cell(row=ws.max_row, column=i), valor, tipo)
            if isinstance(valor, str):
                anchos[i - 1] = min(max(anchos[i - 1], len(valor)), 60)
    for i, ancho in enumerate(anchos, start=1):
        ws.column_dimensions[get_column_letter(i)].width = ancho + 2

    wr = wb.create_sheet("Resumen")
    _encabezados(wr, [f"IVA {direccion} {periodo}", "CFDI", "Pagos", "Base 16 %", "IVA 16 %", "Base 8 %", "IVA 8 %",
                      "Base 0 %", "Base exento", "Base otras tasas", "IVA otras tasas", "Base no objeto", "Retenciones", "IVA total"])

    def _fila(nombre: str, b: dict, negrita: bool = False) -> None:
        wr.append([nombre, b["cfdi"], b["pagos"], b["bases"]["16"], b["iva"]["16"], b["bases"]["8"], b["iva"]["8"],
                   b["bases"]["0"], b["bases"]["exento"], b["bases"]["otras"], b["iva"]["otras"],
                   b["bases"]["no_objeto"], b["retenciones"], b["total"]])
        for celda in wr[wr.max_row][3:]:
            celda.number_format = FORMATO_MONEDA
        if negrita:
            wr[wr.max_row][0].font = _NEGRITA

    bloque = resumen[direccion]
    for clave, etiqueta in _ETIQUETAS_ORIGEN.items():
        _fila(etiqueta, bloque["origenes"][clave])
    _fila(f"Total {direccion}", bloque["total"], negrita=True)
    wr.append([])
    r = resumen["resultado"]
    for etiqueta, clave in (("Trasladado", "trasladado"), ("Acreditable", "acreditable"),
                            ("Retenciones a favor", "retenciones_a_favor"), ("IVA por pagar", "iva_por_pagar"),
                            ("Saldo a cargo", "saldo_a_cargo"), ("Saldo a favor", "saldo_a_favor")):
        wr.append([etiqueta if etiqueta != "IVA por pagar" else "IVA por pagar", float(r[clave])])
        wr.cell(row=wr.max_row, column=2).number_format = FORMATO_MONEDA
    wr.column_dimensions["A"].width = 34
    for i in range(2, 15):
        wr.column_dimensions[get_column_letter(i)].width = 16

    salida = BytesIO()
    wb.save(salida)
    return salida.getvalue()
