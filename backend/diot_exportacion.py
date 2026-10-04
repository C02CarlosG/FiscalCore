"""Excel de la DIOT por flujo: una fila por tercero y tipo de operación, y una hoja de totales y avisos.
Mismas cifras que la pantalla (el cálculo vive en ``diot`` e ``iva_flujo``)."""
from __future__ import annotations

from io import BytesIO

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .diot import ADVERTENCIAS

FORMATO_MONEDA = '"$"#,##0.00'
_ENCABEZADO = Font(bold=True, color="FFFFFF")
_RELLENO = PatternFill("solid", fgColor="334155")

_COLUMNAS = [
    ("RFC", lambda t: t["contraparte_rfc"], "texto"),
    ("Tercero", lambda t: t["contraparte"], "texto"),
    ("Tipo de tercero", lambda t: t["tipo_tercero"], "texto"),
    ("Tipo de operación", lambda t: t["tipo_operacion"], "texto"),
    ("País", lambda t: t["pais"], "texto"),
    ("ID fiscal", lambda t: t["id_fiscal"], "texto"),
    ("CFDI", lambda t: t["cfdi"], "numero"),
    ("Actos 16 %", lambda t: t["actos"]["16"], "moneda"),
    ("Actos 8 %", lambda t: t["actos"]["8"], "moneda"),
    ("Actos 0 %", lambda t: t["actos"]["0"], "moneda"),
    ("Actos exentos", lambda t: t["actos"]["exento"], "moneda"),
    ("Actos otras tasas", lambda t: t["actos"]["otras"], "moneda"),
    ("Actos no objeto", lambda t: t["actos"]["no_objeto"], "moneda"),
    ("IVA 16 %", lambda t: t["iva_pagado"]["16"], "moneda"),
    ("IVA 8 %", lambda t: t["iva_pagado"]["8"], "moneda"),
    ("IVA otras tasas", lambda t: t["iva_pagado"]["otras"], "moneda"),
    ("IVA pagado", lambda t: t["iva_pagado"]["total"], "moneda"),
    ("Devoluciones (valor)", lambda t: t["devoluciones"]["base"], "moneda"),
    ("Devoluciones (IVA)", lambda t: t["devoluciones"]["iva"], "moneda"),
    ("IVA acreditable", lambda t: t["iva_acreditable"], "moneda"),
    ("IVA no acreditable por proporción", lambda t: t["iva_no_acreditable"]["proporcion"], "moneda"),
    ("IVA no acreditable (otros motivos)",
     lambda t: sum((v["iva"] for v in t["iva_no_acreditable"]["por_motivo"].values()), 0), "moneda"),
    ("IVA no acreditable total", lambda t: t["iva_no_acreditable"]["total"], "moneda"),
    ("Retenciones de IVA", lambda t: t["retenciones"], "moneda"),
    ("Advertencias", lambda t: "; ".join(ADVERTENCIAS[a] for a in t["advertencias"]), "texto"),
]


def _encabezados(ws, textos):
    ws.append(textos)
    for celda in ws[1]:
        celda.font = _ENCABEZADO
        celda.fill = _RELLENO
        celda.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"


def construir(resultado: dict) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "DIOT"
    _encabezados(ws, [c[0] for c in _COLUMNAS])
    anchos = [max(len(c[0]), 10) for c in _COLUMNAS]
    for t in resultado["terceros"]:
        fila = [(f(t), tipo) for _, f, tipo in _COLUMNAS]
        ws.append([float(v) if tipo == "moneda" and v is not None else v for v, tipo in fila])
        for i, (celda, (_, tipo)) in enumerate(zip(ws[ws.max_row], fila)):
            if tipo == "moneda":
                celda.number_format = FORMATO_MONEDA
            elif isinstance(celda.value, str):
                anchos[i] = min(max(anchos[i], len(celda.value)), 60)
    for i, ancho in enumerate(anchos, start=1):
        ws.column_dimensions[get_column_letter(i)].width = ancho + 2

    wr = wb.create_sheet("Totales")
    _encabezados(wr, ["Concepto", "Valor"])
    t, c = resultado["totales"], resultado["cuadre_con_iva"]
    for etiqueta, valor, moneda in (
        ("Periodo", resultado["periodo"], False), ("Factor de prorrateo", float(resultado["factor_prorrateo"]), False),
        ("Terceros", t["terceros"], False), ("CFDI", t["cfdi"], False),
        ("Valor de actos", t["valor_de_actos"], True), ("IVA pagado", t["iva_pagado"], True),
        ("Devoluciones (IVA)", t["devoluciones_iva"], True), ("IVA acreditable", t["iva_acreditable"], True),
        ("IVA no acreditable", t["iva_no_acreditable"], True),
        ("IVA acreditable del resumen de IVA", c["iva_acreditable_resumen"], True),
        ("Cuadra con el resumen de IVA", "Sí" if c["cuadra"] else "NO", False),
        ("Terceros con advertencias", t["con_advertencias"], False),
    ):
        wr.append([etiqueta, float(valor) if moneda else valor])
        if moneda:
            wr[wr.max_row][1].number_format = FORMATO_MONEDA
    wr.column_dimensions["A"].width = 38
    wr.column_dimensions["B"].width = 20
    salida = BytesIO()
    wb.save(salida)
    return salida.getvalue()
