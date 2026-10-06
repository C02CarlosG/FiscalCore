"""Router del ISR por flujo (DB mockeada): permisos, validación y delegación."""
import pytest
from fastapi.testclient import TestClient

import backend.main_api as main
from backend import isr_flujo_datos
from backend.deps import get_current_user
from backend.routers import isr_flujo as router

client = TestClient(main.app)
BASE = "/api/v1/empresas/emp-1/isr-flujo"


@pytest.fixture
def con_acceso(monkeypatch):
    main.app.dependency_overrides[get_current_user] = lambda: {"user_id": "u1"}
    monkeypatch.setattr(router, "validar_acceso_empresa", lambda *a, **k: None)
    monkeypatch.setattr(router, "empresa_or_404", lambda eid: {"id": eid, "rfc": "AAA010101AAA", "regimen_fiscal": "612 - PFAE"})
    yield
    main.app.dependency_overrides.clear()


def test_resumen_delega_y_trae_la_aplicabilidad_del_regimen(con_acceso, monkeypatch):
    visto = {}
    monkeypatch.setattr(isr_flujo_datos, "cargar_ajustes", lambda e: {})
    monkeypatch.setattr(isr_flujo_datos, "porcentaje_nomina_exenta", lambda e, ej: visto.update(ejercicio=ej) or isr_flujo_datos.isr_flujo.PORCENTAJE_NOMINA_EXENTA)
    monkeypatch.setattr(isr_flujo_datos, "cargar_eventos", lambda e, rfc, p: visto.update(rfc=rfc, periodo=p) or [])
    monkeypatch.setattr(router, "registrar_evento", lambda *a, **k: None)

    r = client.get(f"{BASE}/2026-09")

    assert r.status_code == 200, r.text
    assert visto == {"ejercicio": 2026, "rfc": "AAA010101AAA", "periodo": "2026-09"}
    d = r.json()
    assert d["regimen"]["codigo"] == "612" and d["regimen"]["modulo"] == "flujo" and d["porcentaje_nomina_exenta"] == 0.47
    assert d["mes"]["utilidad_fiscal_estimada"] == 0.0


@pytest.mark.parametrize("periodo", ["2026-13", "2026-9", "x", "1999-01"])
def test_resumen_rechaza_periodo_invalido(con_acceso, periodo):
    assert client.get(f"{BASE}/{periodo}").status_code == 422


@pytest.mark.parametrize("cuerpo", [{"pct_nomina_exenta": 1.5}, {"pct_nomina_exenta": -0.1}, {"pct_nomina_exenta": 0.5}, {}])
def test_config_rechaza_porcentajes_invalidos(con_acceso, cuerpo):
    assert client.put(f"{BASE}/config/2026", json=cuerpo).status_code == 422


def test_config_guarda_con_el_usuario(con_acceso, monkeypatch):
    visto = {}
    monkeypatch.setattr(isr_flujo_datos, "guardar_porcentaje", lambda e, ej, pct, u, *a: visto.update(ej=ej, pct=str(pct), u=u))
    monkeypatch.setattr(isr_flujo_datos, "parametros_provisional", lambda e, ej: {"ptu_pagada": 0, "perdidas_pendientes": 0})

    r = client.put(f"{BASE}/config/2026", json={"pct_nomina_exenta": 0.53})

    assert r.status_code == 200 and visto == {"ej": 2026, "pct": "0.53", "u": "u1"}


def test_ajuste_exige_motivo_con_texto(con_acceso):
    cuerpo = {"uuid": "U1", "lado": "ingreso"}
    assert client.put(f"{BASE}/ajustes", json={**cuerpo, "motivo": "   "}).status_code == 422
    assert client.put(f"{BASE}/ajustes", json={**cuerpo, "lado": "otro", "motivo": "x"}).status_code == 422


def test_detalle_valida_bloque_y_lado(con_acceso):
    assert client.get(f"{BASE}/2026-09/detalle", params={"lado": "ingreso", "bloque": "otro"}).status_code == 422
    assert client.get(f"{BASE}/2026-09/detalle", params={"lado": "x", "bloque": "contado"}).status_code == 422


def test_exportar_arma_el_excel_con_detalle_y_resumen_y_audita(con_acceso, monkeypatch):
    from io import BytesIO

    import openpyxl

    from backend import isr_flujo
    from backend.tests.test_isr_flujo import doc, todos

    eventos = todos(doc("U1", subtotal=1000))
    auditado = []
    monkeypatch.setattr(isr_flujo_datos, "cargar_ajustes", lambda e: {})
    monkeypatch.setattr(isr_flujo_datos, "porcentaje_nomina_exenta", lambda e, ej: isr_flujo.PORCENTAJE_NOMINA_EXENTA)
    monkeypatch.setattr(isr_flujo_datos, "cargar_eventos", lambda e, rfc, p: eventos)
    monkeypatch.setattr(router, "registrar_evento", lambda *a, **k: auditado.append(k["metadata"]))

    r = client.get(f"{BASE}/2026-09/exportar", params={"lado": "ingreso", "bloque": "contado"})

    assert r.status_code == 200, r.text
    wb = openpyxl.load_workbook(BytesIO(r.content))
    assert wb.sheetnames == ["Detalle", "Resumen"]
    assert wb["Detalle"].max_row == 2 and wb["Detalle"]["I2"].value == 1000
    assert auditado[0]["filas"] == 1


def test_exportar_rechaza_bloque_o_periodo_invalido(con_acceso):
    assert client.get(f"{BASE}/2026-09/exportar", params={"lado": "ingreso", "bloque": "otro"}).status_code == 422
    assert client.get(f"{BASE}/2026-13/exportar", params={"lado": "ingreso", "bloque": "contado"}).status_code == 422


def _preparar_provisional(monkeypatch, regimen):
    monkeypatch.setattr(router, "empresa_or_404", lambda eid: {"id": eid, "rfc": "AAA010101AAA", "regimen_fiscal": regimen})
    monkeypatch.setattr(isr_flujo_datos, "cargar_ajustes", lambda e: {})
    monkeypatch.setattr(isr_flujo_datos, "porcentaje_nomina_exenta", lambda e, ej: isr_flujo_datos.isr_flujo.PORCENTAJE_NOMINA_EXENTA)
    monkeypatch.setattr(isr_flujo_datos, "parametros_provisional", lambda e, ej: {"ptu_pagada": 0, "ptu_mes_pago": None, "perdidas_pendientes": 0, "arrendamiento_periodicidad": "mensual", "deduccion_opcional_35": False})
    monkeypatch.setattr(isr_flujo_datos, "cargar_eventos", lambda e, rfc, p: [])
    monkeypatch.setattr(router, "registrar_evento", lambda *a, **k: None)
    monkeypatch.setattr(router.declaraciones_datos, "pagos_del_ejercicio", lambda e, ej, i: {})


def test_pago_provisional_612_se_calcula_con_la_tarifa(con_acceso, monkeypatch):
    from backend.tests.test_isr_flujo import doc, todos

    _preparar_provisional(monkeypatch, "612 - PFAE")
    monkeypatch.setattr(isr_flujo_datos, "cargar_eventos", lambda e, rfc, p: todos(doc("U1", subtotal=100000, fecha_emision=__import__("datetime").date(2026, 1, 10))))

    d = client.get(f"{BASE}/2026-01/pago-provisional").json()

    assert d["calculado"] is True and d["base_gravable"] == 100000.0 and d["tarifa"]["porcentaje"] == 30.0
    assert d["fuente"]["consultado"] == "2026-10-06"


@pytest.mark.parametrize("regimen,motivo", [("601 - General", "coeficiente"), ("626", "regimen_no_soportado")])
def test_pago_provisional_de_otros_regimenes_no_se_calcula(con_acceso, monkeypatch, regimen, motivo):
    _preparar_provisional(monkeypatch, regimen)

    d = client.get(f"{BASE}/2026-09/pago-provisional").json()

    assert d["calculado"] is False and d["motivo"] == motivo
    assert (d.get("ruta") == "/api/v1/empresas/emp-1/isr-provisional/2026-09") == (motivo == "coeficiente")


def test_pago_provisional_rechaza_periodo_invalido(con_acceso):
    assert client.get(f"{BASE}/2026-13/pago-provisional").status_code == 422


def test_config_acepta_ptu_y_perdidas_no_negativas(con_acceso, monkeypatch):
    visto = {}
    monkeypatch.setattr(isr_flujo_datos, "guardar_porcentaje", lambda e, ej, pct, u, ptu=None, per=None, *a: visto.update(ptu=str(ptu), per=str(per)))
    monkeypatch.setattr(isr_flujo_datos, "parametros_provisional", lambda e, ej: {"ptu_pagada": 0, "perdidas_pendientes": 0})

    assert client.put(f"{BASE}/config/2026", json={"pct_nomina_exenta": 0.47, "ptu_pagada": "100.50", "perdidas_pendientes": 20}).status_code == 200
    assert visto == {"ptu": "100.50", "per": "20"}
    assert client.put(f"{BASE}/config/2026", json={"pct_nomina_exenta": 0.47, "ptu_pagada": -1}).status_code == 422


def test_pago_provisional_606_usa_el_art_116_del_periodo_con_la_opcion_y_el_predial(con_acceso, monkeypatch):
    import datetime

    from backend.tests.test_isr_flujo import doc, todos

    _preparar_provisional(monkeypatch, "606 - Arrendamiento")
    monkeypatch.setattr(isr_flujo_datos, "parametros_provisional", lambda e, ej: {
        "ptu_pagada": 0, "perdidas_pendientes": 0, "arrendamiento_periodicidad": "mensual", "deduccion_opcional_35": True})
    monkeypatch.setattr(isr_flujo_datos, "cargar_eventos", lambda e, rfc, p: todos(
        doc("U1", subtotal=100000, isr_retenido=10000, fecha_emision=datetime.date(2026, 9, 10))))

    d = client.get(f"{BASE}/2026-09/pago-provisional", params={"predial": "1000"}).json()

    assert d["calculado"] is True and d["articulo"] == "116" and d["deduccion_opcional_35"] is True
    assert d["deducciones_usadas"] == 36000.0 and d["impuesto_causado"] == 12936.08
    assert d["isr_retenido_acreditado"] == 10000.0 and d["pago_del_periodo"] == 2936.08


def test_pago_provisional_rechaza_predial_negativo_o_con_mas_de_dos_decimales(con_acceso):
    assert client.get(f"{BASE}/2026-09/pago-provisional", params={"predial": "-1"}).status_code == 422
    assert client.get(f"{BASE}/2026-09/pago-provisional", params={"predial": "1.234"}).status_code == 422


def test_config_acepta_periodicidad_y_opcion_del_35(con_acceso, monkeypatch):
    visto = {}
    monkeypatch.setattr(isr_flujo_datos, "guardar_porcentaje", lambda e, ej, pct, u, ptu=None, per=None, pd=None, op=None, mes=None: visto.update(pd=pd, op=op))
    monkeypatch.setattr(isr_flujo_datos, "parametros_provisional", lambda e, ej: {})

    assert client.put(f"{BASE}/config/2026", json={"pct_nomina_exenta": 0.47, "arrendamiento_periodicidad": "trimestral", "deduccion_opcional_35": True}).status_code == 200
    assert visto == {"pd": "trimestral", "op": True}
    assert client.put(f"{BASE}/config/2026", json={"pct_nomina_exenta": 0.47, "arrendamiento_periodicidad": "anual"}).status_code == 422


def test_pago_provisional_sin_acceso_es_403(monkeypatch):
    from fastapi import HTTPException

    main.app.dependency_overrides[get_current_user] = lambda: {"user_id": "u1"}

    def sin_acceso(*a, **k):
        raise HTTPException(status_code=403, detail="Sin acceso")
    monkeypatch.setattr(router, "validar_acceso_empresa", sin_acceso)
    try:
        assert client.get(f"{BASE}/2026-09/pago-provisional").status_code == 403
    finally:
        main.app.dependency_overrides.clear()


def test_config_acepta_el_mes_de_pago_de_la_ptu(con_acceso, monkeypatch):
    visto = {}
    monkeypatch.setattr(isr_flujo_datos, "guardar_porcentaje",
                        lambda e, ej, pct, u, ptu=None, per=None, pd=None, op=None, mes=None: visto.update(mes=mes))
    monkeypatch.setattr(isr_flujo_datos, "parametros_provisional", lambda e, ej: {})

    assert client.put(f"{BASE}/config/2026", json={"pct_nomina_exenta": 0.47, "ptu_pagada": 100, "ptu_mes_pago": 5}).status_code == 200
    assert visto == {"mes": 5}
    assert client.put(f"{BASE}/config/2026", json={"pct_nomina_exenta": 0.47, "ptu_mes_pago": 13}).status_code == 422
