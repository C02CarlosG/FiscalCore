"""Exportación a Excel del listado de CFDI (DB mockeada): contenido del libro,
límite de filas, columnas del catálogo y auditoría."""
from io import BytesIO

import openpyxl
import pytest
from fastapi.testclient import TestClient

import backend.main_api as main
from backend import auditoria, cfdi_exportacion, cfdi_listado, db
from backend.cfdi_columnas import columnas
from backend.deps import get_current_user
from backend.routers import cfdis

client = TestClient(main.app)
URL = "/api/v1/empresas/emp-1/cfdis/exportar"
OK = {"direccion": "emitidos", "periodo": "2026-09"}
POR_CLAVE = {c.clave: c for c in columnas("emitidos", "I")}
TOTALES = {"conteo": 2, "retencion_iva": 0.0, "retencion_ieps": 0.0, "retencion_isr": 0.0, "traslado_iva": 32.0,
           "traslado_ieps": 0.0, "traslado_isr": 0.0, "total_retenciones": 0.0, "subtotal": 200.0,
           "descuento": 0.0, "neto": 200.0, "total": 232.0}
SIN_DATOS = {k: (0 if k == "conteo" else None) for k in TOTALES}


@pytest.fixture
def con_acceso(monkeypatch):
    main.app.dependency_overrides[get_current_user] = lambda: {"user_id": "u1"}
    monkeypatch.setattr(cfdis, "validar_acceso_empresa", lambda *a, **k: None)
    monkeypatch.setattr(cfdis, "empresa_or_404", lambda eid: {"id": eid, "rfc": "AAA010101AAA"})
    monkeypatch.setattr(cfdi_listado, "resumen",
                        lambda *a: {"conteos": {}, "totales": {"periodo": TOTALES, "acumulado": SIN_DATOS}})
    yield
    main.app.dependency_overrides.clear()


def _libro(contenido: bytes):
    return openpyxl.load_workbook(BytesIO(contenido))


def test_libro_tiene_hoja_de_cfdi_y_de_totales_con_formatos():
    cols = [POR_CLAVE[k] for k in ("fecha_emision", "serie", "total", "pagos_relacionados", "estado")]
    items = [{"fecha_emision": "2026-09-03", "serie": "A", "total": 116.5, "pagos_relacionados": ["U1", "U2"],
              "estado": "vigente"},
             {"fecha_emision": "2026-09-04", "serie": None, "total": 0, "pagos_relacionados": [], "estado": "cancelado"}]
    libro = _libro(cfdi_exportacion.construir(cols, items, {"totales": {"periodo": TOTALES, "acumulado": SIN_DATOS}}))

    assert libro.sheetnames == ["CFDI", "Totales"]
    hoja = libro["CFDI"]
    assert [c.value for c in hoja[1]] == ["Fecha expedición", "Serie", "Total", "CFDIs de pago relacionados", "Estado"]
    assert hoja.max_row == 3
    assert hoja["A2"].value.date().isoformat() == "2026-09-03" and hoja["A2"].number_format == "dd/mm/yyyy"
    assert hoja["C2"].value == 116.5 and "$" in hoja["C2"].number_format
    assert hoja["D2"].value == "U1, U2"
    assert hoja["B3"].value is None and hoja["D3"].value is None   # vacío = celda vacía, no "—"

    totales = libro["Totales"]
    assert [c.value for c in totales["A"]] == [None, "Periodo", "Acumulado"]
    assert totales["M2"].value == 232.0
    assert totales["M3"].value is None   # sin datos: celda vacía, no cero


def test_exporta_con_las_columnas_pedidas_en_su_orden(con_acceso, monkeypatch):
    visto = {}

    def _exportar(empresa_id, rfc, consulta, claves):
        visto.update(claves=claves, orden=consulta.orden, filtros=consulta.filtros)
        return {"columnas": [POR_CLAVE[k] for k in claves], "items": [{"total": 10, "serie": "Z"}], "total": 1}

    monkeypatch.setattr(cfdi_listado, "exportar", _exportar)
    eventos = []
    monkeypatch.setattr(auditoria, "registrar_evento", lambda *a, **k: eventos.append((a, k)))

    r = client.get(URL, params={**OK, "columnas": "total, serie", "orden": "total", "dir": "desc"})

    assert r.status_code == 200
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert 'filename="cfdi_emitidos_I_2026-09.xlsx"' in r.headers["content-disposition"]
    assert visto["claves"] == ["total", "serie"] and visto["orden"] == "total"
    assert [c.value for c in _libro(r.content)["CFDI"][1]] == ["Total", "Serie"]
    (args, kwargs), = eventos
    assert args[:2] == ("u1", "cfdi_exportado") and kwargs["metadata"]["filas"] == 1


def test_sin_columnas_usa_las_visibles_por_defecto(con_acceso, monkeypatch):
    visto = {}

    def _exportar(empresa_id, rfc, consulta, claves):
        visto["claves"] = claves
        return {"columnas": [], "items": [], "total": 0}

    monkeypatch.setattr(cfdi_listado, "exportar", _exportar)
    monkeypatch.setattr(auditoria, "registrar_evento", lambda *a, **k: None)
    assert client.get(URL, params=OK).status_code == 200
    assert visto["claves"] is None


def test_mas_de_50000_filas_responde_422_sin_traer_las_filas(con_acceso, monkeypatch):
    monkeypatch.setattr(db, "query_one", lambda *a, **k: {"n": 50_001})
    monkeypatch.setattr(db, "query_all", lambda *a, **k: pytest.fail("no debe traer las filas"))
    r = client.get(URL, params=OK)
    assert r.status_code == 422
    assert "50,000" in r.json()["detail"]


def test_exactamente_50000_filas_si_se_exporta(con_acceso, monkeypatch):
    monkeypatch.setattr(db, "query_one", lambda *a, **k: {"n": 50_000})
    monkeypatch.setattr(db, "query_all", lambda *a, **k: [])
    monkeypatch.setattr(auditoria, "registrar_evento", lambda *a, **k: None)
    assert client.get(URL, params=OK).status_code == 200


@pytest.mark.parametrize("params", [
    {"columnas": "total,no_existe"},
    {"columnas": "total,total"},
    {"columnas": "total; DROP TABLE cfdi"},
    {"orden": "fecha_emision; DROP TABLE cfdi"},
    {"periodo": "2026-13"},
    {"filtros": "no es json"},
])
def test_parametros_invalidos_responden_422(con_acceso, monkeypatch, params):
    monkeypatch.setattr(db, "query_one", lambda *a, **k: {"n": 1})
    monkeypatch.setattr(db, "query_all", lambda *a, **k: pytest.fail("no debe llegar a las filas"))
    assert client.get(URL, params={**OK, **params}).status_code == 422
