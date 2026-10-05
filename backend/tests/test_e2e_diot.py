"""E2E de la DIOT por flujo contra Postgres real: terceros, clasificación por periodo y por CFDI, cuadre con el IVA, Excel."""
from io import BytesIO

import openpyxl
import pytest

from backend.tests.conftest import db_disponible, headers_usuario_e2e

RFC = "DIO010101E2E"
PROV = "PRO010101AAA"
PROV2 = "OTR010101BBB"
EXT = "XEXX010101000"
EMAIL = "e2e-diot@test.local"
EMAIL_AJENO = "e2e-diot-ajeno@test.local"

pytestmark = [pytest.mark.db, pytest.mark.skipif(not db_disponible(), reason="Postgres no disponible (docker compose up -d db)")]


def _uuid(n):
    return f"8F8A{n:04d}-0000-4000-8000-000000000000"


def _t(base, importe, tasa="0.16"):
    return ("traslado", "002", "Tasa", tasa, base, importe)


def _recibido(db, e, n, impuestos=(), emisor=PROV, nombre="PROVEEDOR UNO", **kw):
    v = dict(uuid=_uuid(n), tipo_comprobante="I", serie="A", folio=str(n), rfc_emisor=emisor, nombre_emisor=nombre,
             rfc_receptor=RFC, nombre_receptor="Empresa DIOT", fecha_emision="2026-09-10 10:00:00", subtotal="1000",
             descuento="0", iva_trasladado="160", iva_retenido="0", isr_retenido="0", total="1160", estado="vigente",
             metodo_pago="PUE", forma_pago="03", uso_cfdi="G03", moneda="MXN", tipo_cambio="1", monto_cobrado="0",
             cfdi_relacionados="[]", es_anticipo_sat=False)
    v.update(kw)
    fila = db.execute(f"INSERT INTO cfdi (empresa_id, {', '.join(v)}) VALUES (%s, {', '.join(['%s'] * len(v))}) RETURNING id",
                      (e, *v.values()), returning=True)
    for ambito, impuesto, factor, tasa, base, importe in impuestos:
        db.execute("INSERT INTO cfdi_impuestos (cfdi_id, ambito, impuesto, tipo_factor, tasa_o_cuota, base, importe)"
                   " VALUES (%s, %s, %s, %s, %s, %s, %s)", (str(fila["id"]), ambito, impuesto, factor, tasa, base, importe))
    return str(fila["id"])


def _limpiar(db):
    db.execute("DELETE FROM empresas WHERE rfc = %s", (RFC,))
    db.execute("DELETE FROM cfdi WHERE uuid LIKE '8F8A%%'")
    db.execute("DELETE FROM usuarios WHERE email IN (%s, %s)", (EMAIL, EMAIL_AJENO))


def _sembrar(db, e):
    _recibido(db, e, 1, [_t(1000, 160)])                                                           # PROV: 1000 + 160
    _recibido(db, e, 2, [_t(500, 80)], subtotal="500", iva_trasladado="80", total="580")           # PROV: 500 + 80
    _recibido(db, e, 3, [_t(100, 16)], tipo_comprobante="E", subtotal="100", iva_trasladado="16", total="116")   # devolución
    _recibido(db, e, 4, [_t(250, 40, "0.16")], emisor=PROV2, nombre="PROVEEDOR DOS", subtotal="250",
              iva_trasladado="40", total="290", uso_cfdi="S01")                                    # uso sin efectos: no acreditable
    _recibido(db, e, 5, [_t(5000, 800)], emisor=PROV2, nombre="PROVEEDOR DOS", subtotal="5000",
              iva_trasladado="800", total="5800", forma_pago="01")                                 # efectivo > 2000: no acreditable
    _recibido(db, e, 6, [_t(300, 48)], emisor=EXT, nombre="ACME INC", subtotal="300", iva_trasladado="48", total="348")
    _recibido(db, e, 7, [_t(400, 64)], emisor=EXT, nombre="GLOBEX LLC", subtotal="400", iva_trasladado="64", total="464")
    _recibido(db, e, 8, [_t(900, 144)], emisor=PROV2, nombre="PROVEEDOR DOS", subtotal="900", iva_trasladado="144",
              total="1044", fecha_emision="2026-08-10 10:00:00")                                   # otro mes: no está en septiembre


@pytest.fixture(scope="module")
def entorno():
    from fastapi.testclient import TestClient

    import backend.main_api as main
    from backend import db

    db.init_db()
    _limpiar(db)
    client = TestClient(main.app)
    try:
        headers = headers_usuario_e2e(db, EMAIL)
        r = client.post("/api/v1/mis-empresas", headers=headers, json={"rfc": RFC, "razon_social": "Empresa DIOT"})
        assert r.status_code == 201, r.text
        e = r.json()["empresa_id"]
        _sembrar(db, e)
        yield db, client, headers, e
    finally:
        _limpiar(db)


def _url(ent, ruta):
    return f"/api/v1/empresas/{ent[3]}/diot-flujo/{ruta}"


def _diot(ent, periodo="2026-09", **params):
    r = ent[1].get(_url(ent, periodo), headers=ent[2], params=params)
    assert r.status_code == 200, r.text
    return r.json()


def _por(d, rfc, nombre=None, operacion=None):
    return next(t for t in d["terceros"] if t["contraparte_rfc"] == rfc and (nombre is None or t["contraparte"] == nombre)
                and (operacion is None or t["tipo_operacion"] == operacion))


def test_terceros_con_valor_de_actos_iva_y_devoluciones(entorno):
    d = _diot(entorno)

    uno = _por(d, PROV)
    assert (uno["cfdi"], uno["actos"]["16"], uno["iva_pagado"]["16"]) == (3, 1500.0, 240.0)
    assert uno["devoluciones"] == {"base": 100.0, "iva": 16.0}
    assert uno["iva_acreditable"] == 224.0
    assert (uno["tipo_tercero"], uno["tipo_operacion"], uno["advertencias"]) == ("04", "85", [])


def test_iva_no_acreditable_por_motivo(entorno):
    dos = _por(_diot(entorno), PROV2)

    assert dos["iva_acreditable"] == 0.0
    por_motivo = dos["iva_no_acreditable"]["por_motivo"]
    assert por_motivo["uso_no_deducible"]["iva"] == 40.0 and por_motivo["efectivo"]["iva"] == 800.0
    assert dos["iva_no_acreditable"]["total"] == 840.0


def test_extranjeros_con_el_rfc_generico_salen_separados_y_pendientes(entorno):
    d = _diot(entorno)

    acme, globex = _por(d, EXT, "ACME INC"), _por(d, EXT, "GLOBEX LLC")
    assert (acme["tipo_tercero"], acme["advertencias"]) == ("05", ["extranjero_pendiente"])
    assert acme["iva_acreditable"] == 48.0 and globex["iva_acreditable"] == 64.0


def test_la_diot_cuadra_con_el_acreditable_del_iva(entorno):
    db, client, headers, e = entorno
    for factor in (1.0, 0.5):
        d = _diot(entorno, factor=factor)
        iva = client.get(f"/api/v1/empresas/{e}/iva-flujo/2026-09", headers=headers, params={"factor": factor}).json()

        assert d["cuadre_con_iva"]["cuadra"] is True
        assert d["totales"]["iva_acreditable"] == iva["acreditable"]["ajustado"]


def test_clasificacion_por_periodo_manda_sobre_el_catalogo_y_se_audita(entorno):
    db, client, headers, e = entorno
    pid = _por(_diot(entorno), PROV)["proveedor_id"]

    r = client.put(_url(entorno, f"2026-09/terceros/{pid}"), headers=headers, json={"tipo_operacion": "06"})

    assert r.status_code == 200, r.text
    assert _por(_diot(entorno), PROV)["tipo_operacion"] == "06"
    assert _por(_diot(entorno, "2026-08"), PROV2)["tipo_operacion"] == "85"                # otro periodo: el catálogo
    aud = db.query_all("SELECT metadata FROM auditoria WHERE empresa_id = %s AND accion = 'diot_tercero_periodo'", (e,))
    assert len(aud) == 1 and aud[0]["metadata"] == {"periodo": "2026-09", "tipo_tercero": None, "tipo_operacion": "06"}
    assert client.delete(_url(entorno, f"2026-09/terceros/{pid}"), headers=headers).status_code == 204
    assert _por(_diot(entorno), PROV)["tipo_operacion"] == "85"


def test_la_clasificacion_valida_las_reglas_cruzadas(entorno):
    client, headers = entorno[1], entorno[2]
    pid = _por(_diot(entorno), PROV)["proveedor_id"]

    assert client.put(_url(entorno, f"2026-09/terceros/{pid}"), headers=headers, json={"tipo_operacion": "87"}).status_code == 422
    assert client.put(_url(entorno, f"2026-09/terceros/{pid}"), headers=headers, json={"tipo_tercero": "05"}).status_code == 422
    assert client.put(_url(entorno, f"2026-09/terceros/{pid}"), headers=headers, json={"tipo_operacion": "99"}).status_code == 422
    assert client.put(_url(entorno, "2026-09/terceros/no-es-uuid"), headers=headers, json={}).status_code == 404


def test_un_tercero_con_dos_operaciones_en_el_mismo_periodo(entorno):
    db, client, headers, e = entorno

    assert client.put(_url(entorno, f"2026-09/cfdi/{_uuid(2)}"), headers=headers, json={"tipo_operacion": "03"}).status_code == 200
    d = _diot(entorno)

    a, b = _por(d, PROV, operacion="85"), _por(d, PROV, operacion="03")
    assert (a["cfdi"], b["cfdi"]) == (2, 1) and b["iva_acreditable"] == 80.0
    assert d["cuadre_con_iva"]["cuadra"] is True and d["operaciones_por_cfdi"] == 1
    assert client.delete(_url(entorno, f"2026-09/cfdi/{_uuid(2)}"), headers=headers).status_code == 204
    assert client.delete(_url(entorno, f"2026-09/cfdi/{_uuid(2)}"), headers=headers).status_code == 404


def test_operacion_de_un_cfdi_que_no_es_recibido_es_404(entorno):
    client, headers = entorno[1], entorno[2]

    assert client.put(_url(entorno, "2026-09/cfdi/no-existe"), headers=headers, json={"tipo_operacion": "03"}).status_code == 404
    assert client.put(_url(entorno, f"2026-09/cfdi/{_uuid(1)}"), headers=headers, json={"tipo_operacion": "99"}).status_code == 422


def test_exportar_excel_trae_los_terceros_y_los_totales(entorno):
    r = entorno[1].get(_url(entorno, "2026-09/exportar"), headers=entorno[2])

    assert r.status_code == 200 and "spreadsheetml" in r.headers["content-type"]
    wb = openpyxl.load_workbook(BytesIO(r.content))
    filas = list(wb["DIOT"].iter_rows(values_only=True))
    assert filas[0][0] == "RFC" and {f[1] for f in filas[1:]} == {"PROVEEDOR UNO", "PROVEEDOR DOS", "ACME INC", "GLOBEX LLC"}
    totales = dict(wb["Totales"].iter_rows(min_row=2, values_only=True))
    assert totales["Cuadra con el resumen de IVA"] == "Sí" and totales["Terceros"] == 4


def test_periodo_y_factor_invalidos_son_422(entorno):
    client, headers = entorno[1], entorno[2]

    assert client.get(_url(entorno, "2026-13"), headers=headers).status_code == 422
    assert client.get(_url(entorno, "2026-09"), headers=headers, params={"factor": 2}).status_code == 422


def test_otra_empresa_recibe_403(entorno):
    db, client = entorno[0], entorno[1]
    db.execute("DELETE FROM usuarios WHERE email = %s", (EMAIL_AJENO,))
    ajeno = headers_usuario_e2e(db, EMAIL_AJENO)

    assert client.get(_url(entorno, "2026-09"), headers=ajeno).status_code == 403
    assert client.get(_url(entorno, "2026-09/exportar"), headers=ajeno).status_code == 403
    assert client.put(_url(entorno, f"2026-09/cfdi/{_uuid(1)}"), headers=ajeno, json={"tipo_operacion": "03"}).status_code == 403


def test_un_extranjero_renombrado_no_se_duplica_al_sincronizar_y_conserva_sus_cfdi(entorno):
    db, client, headers, e = entorno
    acme = _por(_diot(entorno), EXT, "ACME INC")

    r = client.patch(f"/api/v1/empresas/{e}/proveedores/{acme['proveedor_id']}", headers=headers,
                     json={"nombre": "ACME INCORPORATED", "id_fiscal": "12-3456", "pais": "usa"})
    assert r.status_code == 200, r.text
    d = _diot(entorno)

    nombres = sorted(t["contraparte"] for t in d["terceros"] if t["contraparte_rfc"] == EXT)
    assert nombres == ["ACME INC", "GLOBEX LLC"]                                          # el CFDI sigue diciendo ACME INC
    assert _por(d, EXT, "ACME INC")["proveedor_id"] == acme["proveedor_id"]
    assert _por(d, EXT, "ACME INC")["advertencias"] == []
    assert db.query_one("SELECT COUNT(*) AS n FROM proveedores WHERE empresa_id = %s AND rfc = %s", (e, EXT))["n"] == 2


def test_la_operacion_de_un_cfdi_exige_efecto_en_ese_periodo(entorno):
    client, headers = entorno[1], entorno[2]

    # el CFDI 8 es de agosto: en septiembre no tiene efecto
    r = client.put(_url(entorno, f"2026-09/cfdi/{_uuid(8)}"), headers=headers, json={"tipo_operacion": "03"})
    assert r.status_code == 422 and "no tiene efecto" in r.json()["detail"]
    assert client.put(_url(entorno, f"2026-08/cfdi/{_uuid(8)}"), headers=headers, json={"tipo_operacion": "03"}).status_code == 200
    client.delete(_url(entorno, f"2026-08/cfdi/{_uuid(8)}"), headers=headers)


def test_los_actos_incluyen_lo_no_acreditable_y_el_cuadre_es_exacto_con_varios_terceros(entorno):
    d = _diot(entorno, factor=0.7)
    dos = _por(d, PROV2)

    assert dos["actos"]["16"] == 5250.0 and dos["iva_pagado"]["16"] == 840.0           # 250 (S01) + 5,000 (efectivo)
    assert dos["iva_pagado"]["total"] == dos["iva_acreditable"] + dos["iva_no_acreditable"]["total"]
    assert d["cuadre_con_iva"]["cuadra"] is True
    assert [a["codigo"] for a in d["advertencias"]] == []                               # no hay actos a 8 %


def test_los_terceros_sin_actos_positivos_se_advierten(entorno):
    db, e = entorno[0], entorno[3]
    _recibido(db, e, 20, [_t(100, 16)], emisor="NEG010101DDD", nombre="SOLO DEVOLUCION", tipo_comprobante="E",
              subtotal="100", iva_trasladado="16", total="116")
    try:
        t = _por(_diot(entorno), "NEG010101DDD")

        assert "monto_no_positivo" in t["advertencias"] and "sin_catalogo" not in t["advertencias"]
    finally:
        db.execute("DELETE FROM cfdi WHERE uuid = %s", (_uuid(20),))
