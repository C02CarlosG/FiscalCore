"""Excel del IVA por flujo: hoja de detalle con las bases por tasa y hoja de resumen del periodo. Sin DB."""
from datetime import date
from decimal import Decimal
from io import BytesIO

import openpyxl

from backend import iva_flujo as f
from backend import iva_flujo_exportacion as x
from backend.tests.test_iva_flujo import D, RFC, _construir, doc, imp, pago, recibido, tras


def _libro(contenido: bytes):
    return openpyxl.load_workbook(BytesIO(contenido))


def _datos(docs, pagos=(), direccion="trasladado", origen="contado", ajustes=None, periodo="2026-09"):
    ajustes = ajustes or {}
    eventos = _construir(docs, pagos)
    filas = f.renglones(eventos, periodo, direccion, origen, ajustes)
    resumen = f.resumen(eventos, periodo, ajustes)
    return x.construir(periodo, direccion, origen, filas, resumen)


def test_hojas_y_encabezados():
    wb = _libro(_datos([doc("U1")]))

    assert wb.sheetnames == ["Detalle", "Resumen"]
    encabezados = [c.value for c in wb["Detalle"][1]]
    assert encabezados[:4] == ["Fecha de emisión", "Fecha de pago", "UUID", "UUID del REP"]
    assert "Base 16 %" in encabezados and "IVA 16 %" in encabezados and "IVA total" in encabezados
    assert wb["Detalle"].freeze_panes == "A2"


def test_una_fila_por_cfdi_con_sus_bases_por_tasa():
    wb = _libro(_datos([doc("U1", impuestos=[tras("0.16", 1000, 160), tras("0.08", 500, 40)], iva_trasladado=D("200"))]))
    ws = wb["Detalle"]
    enc = [c.value for c in ws[1]]
    fila = {enc[i]: c.value for i, c in enumerate(ws[2])}

    assert ws.max_row == 2
    assert fila["UUID"] == "U1" and fila["RFC contraparte"] == "XAXX010101000"
    assert fila["Base 16 %"] == 1000 and fila["IVA 16 %"] == 160 and fila["Base 8 %"] == 500 and fila["IVA 8 %"] == 40
    assert fila["IVA total"] == 200 and fila["Fecha de emisión"].date() == date(2026, 9, 10)
    assert fila["Fecha de pago"] is None


def test_credito_trae_fecha_de_pago_y_uuid_del_rep():
    d = doc("U1", metodo_pago="PPD", total=D("1160"))
    ws = _libro(_datos([d], [pago("U1")], origen="credito"))["Detalle"]
    enc = [c.value for c in ws[1]]
    fila = {enc[i]: c.value for i, c in enumerate(ws[2])}

    assert fila["Fecha de pago"].date() == date(2026, 9, 25) and fila["UUID del REP"] == "REP1" and fila["Parcialidad"] == 1


def test_marcas_motivo_y_ajuste_se_exportan_como_texto():
    aj = {("U1", "trasladado"): {"accion": "excluir", "periodo_destino": None, "motivo": "duplicado"}}
    ws = _libro(_datos([doc("U1", iva_trasladado=D("161"))], origen="no_considerados", ajustes=aj))["Detalle"]
    enc = [c.value for c in ws[1]]
    fila = {enc[i]: c.value for i, c in enumerate(ws[2])}

    assert fila["Motivo"] == "manual"
    assert fila["Ajuste"] == "excluir: duplicado"
    assert "descuadre" in fila["Marcas"]


def test_importes_con_formato_de_moneda():
    ws = _libro(_datos([doc("U1")]))["Detalle"]
    enc = [c.value for c in ws[1]]

    assert ws.cell(row=2, column=enc.index("IVA total") + 1).number_format == '"$"#,##0.00'
    assert ws.cell(row=2, column=1).number_format == "dd/mm/yyyy"


def test_sin_renglones_deja_solo_el_encabezado():
    ws = _libro(_datos([], origen="contado"))["Detalle"]

    assert ws.max_row == 1


def test_resumen_trae_cada_origen_y_el_resultado():
    docs = [doc("U1"), recibido("R1")]
    ws = _libro(_datos(docs))["Resumen"]
    valores = {fila[0].value: [c.value for c in fila[1:]] for fila in ws.iter_rows(min_row=2)}

    assert valores["Contado"][-1] == 160                           # IVA total del origen
    assert valores["Total trasladado"][-1] == 160
    assert valores["IVA por pagar"][0] == 0                         # 160 − 160 − 0
    assert "Periodo" in {fila[0].value for fila in ws.iter_rows(min_row=1, max_row=1)} or ws["A1"].value


def test_un_texto_que_parece_formula_no_se_ejecuta():
    d = doc("U1", nombre_receptor="=HYPERLINK(\"http://x\",\"x\")")
    ws = _libro(_datos([d]))["Detalle"]
    enc = [c.value for c in ws[1]]
    celda = ws.cell(row=2, column=enc.index("Contraparte") + 1)

    assert celda.data_type == "s" and str(celda.value).startswith("=")       # texto, no fórmula
