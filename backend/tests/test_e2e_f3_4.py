"""E2E de F3.4 contra Postgres real: exportación a Excel (mismas filas, columnas y
totales que el listado) y preferencias de columnas por usuario. Se salta sin DB."""
from io import BytesIO

import openpyxl
import pytest

from backend.tests.conftest import db_disponible, headers_usuario_e2e
from backend.tests.test_e2e_listado_cfdi import EMAIL_AJENO, entorno  # noqa: F401  (fixture del módulo)

pytestmark = [pytest.mark.db, pytest.mark.skipif(not db_disponible(), reason="Postgres no disponible (docker compose up -d db)")]


def _base(entorno):
    _db, client, headers, empresa_id = entorno
    return client, headers, f"/api/v1/empresas/{empresa_id}/cfdis"


def _exportar(entorno, **params):
    client, headers, base = _base(entorno)
    return client.get(f"{base}/exportar", headers=headers,
                      params={"direccion": "emitidos", "periodo": "2026-03", **params})


def _filas(contenido):
    hoja = openpyxl.load_workbook(BytesIO(contenido))["CFDI"]
    return [[c.value for c in fila] for fila in hoja.iter_rows()]


def test_exporta_todas_las_paginas_con_las_columnas_pedidas_en_orden(entorno):
    client, headers, base = _base(entorno)
    listado_p1 = client.get(base, headers=headers, params={"direccion": "emitidos", "periodo": "2026-03"}).json()

    r = _exportar(entorno, columnas="total,folio,estado", orden="total", dir="desc")

    assert r.status_code == 200
    filas = _filas(r.content)
    assert filas[0] == ["Total", "Folio", "Estado"]
    assert len(filas) - 1 == listado_p1["total"] == 38          # todas las páginas, no solo la primera
    assert [f[0] for f in filas[1:4]] == [4060.0, 3944.0, 3828.0]
    assert {f[2] for f in filas[1:]} == {"vigente"}


def test_filtros_y_orden_de_la_pantalla_se_respetan(entorno):
    r = _exportar(entorno, columnas="folio,estado", estado="cancelado")
    assert [f[1] for f in _filas(r.content)[1:]] == ["cancelado"]

    r = _exportar(entorno, columnas="folio", filtros='[{"campo": "total", "op": "mayor", "valor": 4000}]')
    assert len(_filas(r.content)) - 1 == 1


def test_totales_del_excel_coinciden_con_los_de_pantalla(entorno):
    client, headers, base = _base(entorno)
    resumen = client.get(f"{base}/resumen", headers=headers,
                         params={"direccion": "emitidos", "periodo": "2026-03"}).json()["totales"]

    hoja = openpyxl.load_workbook(BytesIO(_exportar(entorno).content))["Totales"]

    encabezados = [c.value for c in hoja[1]]
    periodo = dict(zip(encabezados[1:], [c.value for c in hoja[2]][1:]))
    acumulado = dict(zip(encabezados[1:], [c.value for c in hoja[3]][1:]))
    assert periodo["CFDI"] == resumen["periodo"]["conteo"]
    assert periodo["Total"] == resumen["periodo"]["total"] and periodo["Neto"] == resumen["periodo"]["neto"]
    assert acumulado["Total"] == resumen["acumulado"]["total"]


def test_sin_columnas_exporta_las_visibles_por_defecto(entorno):
    client, headers, base = _base(entorno)
    visibles = [c["etiqueta"] for c in client.get(f"{base}/columnas", headers=headers,
                params={"direccion": "emitidos", "tipo": "I"}).json()["encabezado"] if c["visible_por_defecto"]]
    assert _filas(_exportar(entorno).content)[0] == visibles


def test_exportar_deja_rastro_en_auditoria(entorno):
    db = entorno[0]
    antes = db.query_one("SELECT COUNT(*) AS n FROM auditoria WHERE accion = 'cfdi_exportado'")["n"]
    _exportar(entorno, columnas="total")
    assert db.query_one("SELECT COUNT(*) AS n FROM auditoria WHERE accion = 'cfdi_exportado'")["n"] == antes + 1


def test_usuario_sin_acceso_no_puede_exportar(entorno):
    db = entorno[0]
    db.execute("DELETE FROM usuarios WHERE email = %s", (EMAIL_AJENO,))
    ajeno = headers_usuario_e2e(db, EMAIL_AJENO)
    client, _headers, base = _base(entorno)
    r = client.get(f"{base}/exportar", headers=ajeno, params={"direccion": "emitidos", "periodo": "2026-03"})
    assert r.status_code == 403


def test_preferencias_se_guardan_por_usuario_y_se_restablecen(entorno):
    db, client, headers, _empresa = entorno
    url = "/api/v1/preferencias/tablas/cfdi-emitidos-I"
    columnas = [{"clave": "total", "visible": True}, {"clave": "uuid", "visible": False}]

    assert client.get(url, headers=headers).json() == {"columnas": None}
    assert client.put(url, headers=headers, json={"columnas": columnas}).status_code == 200
    assert client.get(url, headers=headers).json() == {"columnas": columnas}

    nuevas = list(reversed(columnas))
    client.put(url, headers=headers, json={"columnas": nuevas})
    assert client.get(url, headers=headers).json() == {"columnas": nuevas}      # reemplaza, no acumula

    db.execute("DELETE FROM usuarios WHERE email = %s", (EMAIL_AJENO,))
    otro = headers_usuario_e2e(db, EMAIL_AJENO)
    assert client.get(url, headers=otro).json() == {"columnas": None}           # ajeno al otro usuario

    assert client.delete(url, headers=headers).status_code == 204
    assert client.get(url, headers=headers).json() == {"columnas": None}
